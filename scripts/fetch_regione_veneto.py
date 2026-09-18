"""Download the official Veneto tourism datasets without altering their contents."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from urllib.parse import urlencode, urljoin

import requests
from openpyxl import load_workbook


BASE_URL = "https://statistica.regione.veneto.it/"
MOVEMENT_CSV_URL = urljoin(BASE_URL, "jsp/turi1txt.jsp")
COMMUNE_CSV_URL = urljoin(BASE_URL, "jsp/turismo_comune6txt.jsp")
HOMEPAGE_URL = BASE_URL

TERRITORIES = {
    "provincia_belluno": {
        "D1": "PROVINCIA",
        "D2": "02Belluno",
        "D3": "Movimento annuale per mese",
    },
    "stl_dolomiti": {
        "D1": "STL",
        "D2": "0001 Dolomiti",
        "D3": "Movimento annuale per mese",
    },
    "stl_belluno": {
        "D1": "STL",
        "D2": "0102 Belluno",
        "D3": "Movimento annuale per mese",
    },
}

MONTH_NAMES = {
    "gennaio": 1,
    "febbraio": 2,
    "marzo": 3,
    "aprile": 4,
    "maggio": 5,
    "giugno": 6,
    "luglio": 7,
    "agosto": 8,
    "settembre": 9,
    "ottobre": 10,
    "novembre": 11,
    "dicembre": 12,
}

USER_AGENT = (
    "DMO-Dolomiti-Bellunesi-data-pipeline/1.0 "
    "(+https://github.com/stefanocasagrande-hash/dmodolomiti-turismo-veneto)"
)


class DownloadError(RuntimeError):
    """Raised when a source cannot be downloaded or has an unexpected format."""


def _get(url: str, *, timeout: int = 60, attempts: int = 3) -> bytes:
    last_error: Exception | None = None
    for attempt in range(1, attempts + 1):
        try:
            response = requests.get(
                url,
                headers={"User-Agent": USER_AGENT},
                timeout=timeout,
            )
            response.raise_for_status()
            return response.content
        except requests.RequestException as exc:
            last_error = exc
            if attempt < attempts:
                time.sleep(2**attempt)
    raise DownloadError(f"Download fallito dopo {attempts} tentativi: {url}") from last_error


def _sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _write_if_changed(path: Path, content: bytes) -> bool:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and path.read_bytes() == content:
        return False
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_bytes(content)
    temporary.replace(path)
    return True


def _csv_url(base: str, params: dict[str, object]) -> str:
    return f"{base}?{urlencode(params)}"


def build_year_sources(year: int) -> list[tuple[str, str]]:
    sources: list[tuple[str, str]] = []
    for slug, params in TERRITORIES.items():
        query = {"D0": year, **params, "B1": "Visualizza"}
        sources.append((f"{slug}.csv", _csv_url(MOVEMENT_CSV_URL, query)))

    for metric, x1 in (("arrivi", 2), ("presenze", 3)):
        query = {"anno": year, "provenienza": 0, "x1": x1}
        sources.append(
            (f"comuni_veneto_{metric}.csv", _csv_url(COMMUNE_CSV_URL, query))
        )
    return sources


def _validate_download(name: str, content: bytes) -> None:
    if name.endswith(".csv"):
        first_line = content.splitlines()[0].lower() if content else b""
        if b"progressivo" not in first_line:
            raise DownloadError(f"Il file {name} non contiene l'intestazione CSV attesa")
    elif name.endswith(".xlsx"):
        if not content.startswith(b"PK"):
            raise DownloadError(f"Il file {name} non e' un workbook XLSX valido")
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise DownloadError(f"Il workbook {name} non e' leggibile") from exc
        required_sheets = {"Province", "STL"}
        if not required_sheets.issubset(workbook.sheetnames):
            raise DownloadError(
                f"Il workbook {name} non contiene i fogli attesi: {sorted(required_sheets)}"
            )


def _load_manifest(path: Path) -> dict:
    if not path.exists():
        return {"files": {}}
    return json.loads(path.read_text(encoding="utf-8"))


def _save_manifest(path: Path, manifest: dict) -> None:
    serialized = json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    if not path.exists() or path.read_text(encoding="utf-8") != serialized:
        path.write_text(serialized, encoding="utf-8")


def _download_one(raw_root: Path, year: int, name: str, url: str) -> dict:
    content = _get(url)
    _validate_download(name, content)
    destination = raw_root / str(year) / name
    changed = _write_if_changed(destination, content)
    return {
        "relative_path": destination.relative_to(raw_root).as_posix(),
        "source_url": url,
        "sha256": _sha256(content),
        "bytes": len(content),
        "changed": changed,
    }


def discover_latest_summary_url() -> str:
    homepage = _get(HOMEPAGE_URL).decode("cp1252", errors="replace")
    notice_links = re.findall(
        r'href=["\']([^"\']*novita/notizia_[^"\']+\.jsp)["\']',
        homepage,
        flags=re.IGNORECASE,
    )
    seen: set[str] = set()
    for link in notice_links[:16]:
        notice_url = urljoin(HOMEPAGE_URL, link)
        if notice_url in seen:
            continue
        seen.add(notice_url)
        notice = _get(notice_url).decode("cp1252", errors="replace")
        matches = re.findall(
            r'href=["\']([^"\']*Movimento_turistico_[^"\']+\.xlsx)["\']',
            notice,
            flags=re.IGNORECASE,
        )
        if matches:
            return urljoin(notice_url, matches[0])
    raise DownloadError("Nessun Excel mensile del movimento turistico trovato nelle novita'")


def _summary_sort_key(url: str) -> tuple[int, int]:
    filename = url.rsplit("/", 1)[-1].lower()
    year_match = re.search(r"(20\d{2})\.xlsx", filename)
    month_match = re.search(r"gennaio_([a-z]+)_20\d{2}", filename)
    year = int(year_match.group(1)) if year_match else 0
    month = MONTH_NAMES.get(month_match.group(1), 1) if month_match else 1
    return year, month


def fetch_official_data(
    years: Iterable[int],
    raw_root: Path,
    *,
    include_summary: bool = True,
    max_workers: int = 4,
) -> dict:
    raw_root.mkdir(parents=True, exist_ok=True)
    manifest_path = raw_root / "manifest.json"
    manifest = _load_manifest(manifest_path)
    previous_files = manifest.get("files", {})

    jobs = [
        (year, name, url)
        for year in sorted(set(years))
        for name, url in build_year_sources(year)
    ]
    results: list[dict] = []
    with ThreadPoolExecutor(max_workers=max_workers) as pool:
        futures = {
            pool.submit(_download_one, raw_root, year, name, url): (year, name)
            for year, name, url in jobs
        }
        for future in as_completed(futures):
            results.append(future.result())

    if include_summary:
        summary_url = discover_latest_summary_url()
        year, _ = _summary_sort_key(summary_url)
        name = summary_url.rsplit("/", 1)[-1]
        results.append(_download_one(raw_root, year, name, summary_url))

    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    files: dict[str, dict] = dict(previous_files)
    for result in results:
        path = result["relative_path"]
        previous = previous_files.get(path, {})
        acquired_at = (
            now
            if result["changed"] or previous.get("sha256") != result["sha256"]
            else previous.get("acquired_at_utc", now)
        )
        files[path] = {
            "acquired_at_utc": acquired_at,
            "bytes": result["bytes"],
            "sha256": result["sha256"],
            "source_url": result["source_url"],
        }

    manifest = {
        "description": "Archivio delle fonti ufficiali Regione Veneto, conservate senza modifiche.",
        "files": files,
        "source": "Regione del Veneto - U.O. Sistema Statistico Regionale",
    }
    _save_manifest(manifest_path, manifest)
    return {
        "downloaded": len(results),
        "changed": sum(bool(item["changed"]) for item in results),
        "manifest": str(manifest_path),
    }


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--years", nargs="+", type=int, required=True)
    parser.add_argument(
        "--raw-root", type=Path, default=Path("data/raw/regione_veneto")
    )
    parser.add_argument("--skip-summary", action="store_true")
    parser.add_argument("--max-workers", type=int, default=4)
    return parser.parse_args()


def main() -> None:
    args = _parse_args()
    result = fetch_official_data(
        args.years,
        args.raw_root,
        include_summary=not args.skip_summary,
        max_workers=args.max_workers,
    )
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
